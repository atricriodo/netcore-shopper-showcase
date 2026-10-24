#!/usr/bin/env python3
"""
build_presentation.py
Generates an executive, highly-polished, 16:9 widescreen PowerPoint presentation (.pptx)
from the Netcore unified commerce landing page.
Preserves real product imagery, exact brand colors, structured bento grids, and clear typography.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ================= Brand Color System =================
COLOR_BLAZE       = RGBColor(252, 92, 5)     # #FC5C05 Primary Blaze Orange
COLOR_BLAZE_LIGHT = RGBColor(255, 138, 61)   # #FF8A3D Blaze Highlight
COLOR_BASTILLE    = RGBColor(40, 32, 44)     # #28202C Deep Bastille
COLOR_NIGHT       = RGBColor(26, 20, 30)     # #1A141E Darker Night
COLOR_CARD_DARK   = RGBColor(48, 39, 53)     # #302735 Dark Card Surface
COLOR_BORDER_DARK = RGBColor(75, 62, 82)     # #4B3E52 Dark Card Border
COLOR_WHITE       = RGBColor(255, 255, 255)  # #FFFFFF Pure White
COLOR_MIST        = RGBColor(248, 247, 250)  # #F8F7FA Light Page Background
COLOR_CARD_LIGHT  = RGBColor(255, 255, 255)  # Light Card Surface
COLOR_BORDER_LIGHT= RGBColor(226, 222, 230)  # Light Card Border
COLOR_PEACH       = RGBColor(255, 237, 226)  # #FFEDE2 Peach soft
COLOR_PEACH_BORDER= RGBColor(255, 217, 194)  # #FFD9C2
COLOR_GREEN       = RGBColor(31, 163, 106)   # #1FA36A Success / Live
COLOR_GREEN_SOFT  = RGBColor(228, 246, 237)  # Mint soft
COLOR_RED         = RGBColor(226, 54, 54)    # Warning / Miss
COLOR_RED_SOFT    = RGBColor(254, 238, 238)  # Warning soft
COLOR_INK         = RGBColor(33, 26, 37)     # #211A25 Body text dark
COLOR_MUTED       = RGBColor(110, 103, 116)  # #6E6774 Secondary text
COLOR_FAINT       = RGBColor(165, 158, 170)  # #A59EAA Muted on dark

FONT_FAMILY = "Arial"

def add_full_bg(slide, prs, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_card(slide, left, top, width, height, bg_color, border_color=None, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_header(slide, tag_text, title_text, subtitle_text=None, dark_mode=False):
    tag_bg = COLOR_CARD_DARK if dark_mode else COLOR_PEACH
    tag_border = COLOR_BLAZE if dark_mode else COLOR_PEACH_BORDER
    tag_card = add_card(slide, Inches(0.8), Inches(0.48), Inches(2.8), Inches(0.32), tag_bg, tag_border, 1)
    tf_tag = tag_card.text_frame
    tf_tag.word_wrap = False
    tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = FONT_FAMILY
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_BLAZE
    p_tag.alignment = PP_ALIGN.CENTER

    brand_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.42), Inches(2.0), Inches(0.4))
    tf_brand = brand_box.text_frame
    tf_brand.margin_left = tf_brand.margin_right = tf_brand.margin_top = tf_brand.margin_bottom = 0
    p_br = tf_brand.paragraphs[0]
    p_br.alignment = PP_ALIGN.RIGHT
    run1 = p_br.add_run()
    run1.text = "netcore"
    run1.font.name = FONT_FAMILY
    run1.font.size = Pt(13)
    run1.font.bold = True
    run1.font.color.rgb = COLOR_WHITE if dark_mode else COLOR_INK
    run2 = p_br.add_run()
    run2.text = " •"
    run2.font.name = FONT_FAMILY
    run2.font.size = Pt(13)
    run2.font.bold = True
    run2.font.color.rgb = COLOR_BLAZE

    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(1.0))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p_title = tf.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_FAMILY
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE if dark_mode else COLOR_INK

    if subtitle_text:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = COLOR_FAINT if dark_mode else COLOR_MUTED
        p_sub.space_before = Pt(4)

def safe_add_image(slide, img_path, left, top, width, height):
    if os.path.exists(img_path):
        return slide.shapes.add_picture(img_path, left, top, width, height)
    print(f"Warning: Image {img_path} not found.")
    return None

# ================= SLIDE BUILDERS =================

def build_slide_1_hero(prs):
    """Slide 1: Executive Title & Overview (Dark Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_BASTILLE)

    brand_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(3.0), Inches(0.5))
    tf_b = brand_box.text_frame
    p_b = tf_b.paragraphs[0]
    r1 = p_b.add_run()
    r1.text = "netcore"
    r1.font.name = FONT_FAMILY
    r1.font.size = Pt(20)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_WHITE
    r2 = p_b.add_run()
    r2.text = " •"
    r2.font.name = FONT_FAMILY
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_BLAZE

    pill = add_card(slide, Inches(0.8), Inches(1.3), Inches(3.4), Inches(0.36), COLOR_CARD_DARK, COLOR_BLAZE)
    p_p = pill.text_frame.paragraphs[0]
    p_p.text = "✨ UNIFIED COMMERCE EXPERIENCE PLATFORM"
    p_p.font.name = FONT_FAMILY
    p_p.font.size = Pt(9.5)
    p_p.font.bold = True
    p_p.font.color.rgb = COLOR_BLAZE_LIGHT
    p_p.alignment = PP_ALIGN.CENTER

    copy_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(4.2))
    tf_c = copy_box.text_frame
    tf_c.word_wrap = True

    p1 = tf_c.paragraphs[0]
    p1.text = "Market to Every Shopper Like You "
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    r_hl = p1.add_run()
    r_hl.text = "Already Know Them"
    r_hl.font.color.rgb = COLOR_BLAZE_LIGHT

    p2 = tf_c.add_paragraph()
    p2.text = (
        "Netcore connects on-site search and real-time product discovery directly with "
        "email, SMS, mobile push, and WhatsApp. What a shopper searches, filters, and leaves "
        "in their cart immediately shapes the very next message they receive."
    )
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_FAINT
    p2.space_before = Pt(14)

    bullets = [
        ("⚡ Zero Data Lag", "Real-time intent captured from in-session search & browse behavior."),
        ("🎯 1:1 Relevance", "Automated outbound triggers matching exact size, color & budget."),
        ("🤖 Autonomous AI", "Unified agents driving search ranking and omnichannel campaigns.")
    ]
    for title, desc in bullets:
        pb = tf_c.add_paragraph()
        pb.space_before = Pt(8)
        rb_t = pb.add_run()
        rb_t.text = f"{title}: "
        rb_t.font.name = FONT_FAMILY
        rb_t.font.size = Pt(10.5)
        rb_t.font.bold = True
        rb_t.font.color.rgb = COLOR_WHITE
        rb_d = pb.add_run()
        rb_d.text = desc
        rb_d.font.name = FONT_FAMILY
        rb_d.font.size = Pt(10.5)
        rb_d.font.color.rgb = COLOR_FAINT

    # Right Showcase Card: The Connected System
    showcase = add_card(slide, Inches(6.9), Inches(1.3), Inches(5.6), Inches(4.8), COLOR_CARD_DARK, COLOR_BORDER_DARK)
    
    # 1. On-Site Search Simulation
    s1 = add_card(slide, Inches(7.15), Inches(1.5), Inches(5.1), Inches(1.35), COLOR_BASTILLE, COLOR_BORDER_DARK)
    tb_s1 = slide.shapes.add_textbox(Inches(7.3), Inches(1.6), Inches(3.6), Inches(1.15))
    tf_s1 = tb_s1.text_frame
    tf_s1.word_wrap = True
    p_s1_h = tf_s1.paragraphs[0]
    p_s1_h.text = "ON-SITE STORE SEARCH • 2:02 PM"
    p_s1_h.font.size = Pt(9)
    p_s1_h.font.bold = True
    p_s1_h.font.color.rgb = COLOR_FAINT

    p_s1_q = tf_s1.add_paragraph()
    p_s1_q.text = "🔍 “linen shirt for a summer wedding”"
    p_s1_q.font.size = Pt(11.5)
    p_s1_q.font.bold = True
    p_s1_q.font.color.rgb = COLOR_WHITE
    p_s1_q.space_before = Pt(4)

    p_s1_res = tf_s1.add_paragraph()
    p_s1_res.text = "Filters: Sky Blue • Size M • $60–$90"
    p_s1_res.font.size = Pt(9)
    p_s1_res.font.color.rgb = COLOR_BLAZE_LIGHT
    p_s1_res.space_before = Pt(4)

    safe_add_image(slide, "images/shirt-sky.jpg", Inches(11.05), Inches(1.6), Inches(1.0), Inches(1.0))

    # 2. Netcore Shared Shopper Memory
    s2 = add_card(slide, Inches(7.15), Inches(3.0), Inches(5.1), Inches(1.15), COLOR_NIGHT, COLOR_BLAZE)
    tb_s2 = slide.shapes.add_textbox(Inches(7.3), Inches(3.05), Inches(4.8), Inches(1.0))
    tf_s2 = tb_s2.text_frame
    p_s2_h = tf_s2.paragraphs[0]
    r_live = p_s2_h.add_run()
    r_live.text = "NETCORE SHOPPER MEMORY "
    r_live.font.size = Pt(9.5)
    r_live.font.bold = True
    r_live.font.color.rgb = COLOR_BLAZE
    r_badge = p_s2_h.add_run()
    r_badge.text = "[LIVE SYNC]"
    r_badge.font.size = Pt(8.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = COLOR_GREEN

    p_s2_tags = tf_s2.add_paragraph()
    p_s2_tags.text = "Tags: [linen]  [sky blue]  [size: M]  [affinity: wedding]  [cart: $79]"
    p_s2_tags.font.size = Pt(9.5)
    p_s2_tags.font.bold = True
    p_s2_tags.font.color.rgb = COLOR_WHITE
    p_s2_tags.space_before = Pt(4)

    p_s2_sub = tf_s2.add_paragraph()
    p_s2_sub.text = "Action: Trigger contextual cart recovery + matching chinos"
    p_s2_sub.font.size = Pt(8.5)
    p_s2_sub.font.color.rgb = COLOR_FAINT
    p_s2_sub.space_before = Pt(2)

    # 3. Triggered Outbound Email
    s3 = add_card(slide, Inches(7.15), Inches(4.3), Inches(5.1), Inches(1.6), COLOR_BASTILLE, COLOR_BORDER_DARK)
    tb_s3 = slide.shapes.add_textbox(Inches(7.3), Inches(4.35), Inches(3.0), Inches(1.5))
    tf_s3 = tb_s3.text_frame
    tf_s3.word_wrap = True
    p_s3_h = tf_s3.paragraphs[0]
    p_s3_h.text = "OUTBOUND TRIGGER • EMAIL & WHATSAPP"
    p_s3_h.font.size = Pt(8.5)
    p_s3_h.font.bold = True
    p_s3_h.font.color.rgb = COLOR_GREEN

    p_s3_sub = tf_s3.add_paragraph()
    p_s3_sub.text = "“Your sky blue linen shirt is still in your cart”"
    p_s3_sub.font.size = Pt(11)
    p_s3_sub.font.bold = True
    p_s3_sub.font.color.rgb = COLOR_WHITE
    p_s3_sub.space_before = Pt(3)

    p_s3_body = tf_s3.add_paragraph()
    p_s3_body.text = "Size M in stock. Wedding planning? Pair it with stone chinos."
    p_s3_body.font.size = Pt(9)
    p_s3_body.font.color.rgb = COLOR_PEACH
    p_s3_body.space_before = Pt(2)

    safe_add_image(slide, "images/shirt-sky.jpg", Inches(10.45), Inches(4.6), Inches(0.85), Inches(0.85))
    safe_add_image(slide, "images/pants-stone.jpg", Inches(11.35), Inches(4.6), Inches(0.85), Inches(0.85))

    sp_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
    tf_sp = sp_box.text_frame
    p_sp = tf_sp.paragraphs[0]
    r_sp_h = p_sp.add_run()
    r_sp_h.text = "6,500+ ENTERPRISE BRANDS GROW WITH NETCORE:   "
    r_sp_h.font.size = Pt(9)
    r_sp_h.font.bold = True
    r_sp_h.font.color.rgb = COLOR_MUTED
    r_sp_b = p_sp.add_run()
    r_sp_b.text = "adidas   •   UNDER ARMOUR   •   Columbia   •   crocs   •   ALDO   •   TOMMY HILFIGER   •   Walmart"
    r_sp_b.font.size = Pt(10)
    r_sp_b.font.bold = True
    r_sp_b.font.color.rgb = COLOR_FAINT

def build_slide_2_problem(prs):
    """Slide 2: The Core Problem - The Disconnected Stack (Light Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_MIST)

    add_header(
        slide,
        tag_text="The Status Quo Challenge",
        title_text="Most Stacks Would Have Sent This Instead",
        subtitle_text="When on-site search and customer engagement live in separate tools, campaigns ignore the shopper's real intent.",
        dark_mode=False
    )

    left_card = add_card(slide, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.9), COLOR_WHITE, COLOR_RED, border_width=1.5)
    
    tb_l = slide.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(5.2), Inches(4.5))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p_badge = tf_l.paragraphs[0]
    p_badge.text = "⚠️ THE TRADITIONAL CART ABANDONMENT EMAIL"
    p_badge.font.size = Pt(9.5)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_RED

    p_subj = tf_l.add_paragraph()
    p_subj.text = "“You left something behind!”"
    p_subj.font.size = Pt(15)
    p_subj.font.bold = True
    p_subj.font.color.rgb = COLOR_INK
    p_subj.space_before = Pt(6)

    p_subtext = tf_l.add_paragraph()
    p_subtext.text = (
        "“Still thinking it over? Your cart is waiting. While you’re here, check out this week’s bestsellers:”"
    )
    p_subtext.font.size = Pt(10)
    p_subtext.font.color.rgb = COLOR_MUTED
    p_subtext.space_before = Pt(4)

    safe_add_image(slide, "images/shoe-grey.jpg", Inches(1.1), Inches(3.4), Inches(1.3), Inches(1.3))
    safe_add_image(slide, "images/shirt-grey.jpg", Inches(2.7), Inches(3.4), Inches(1.3), Inches(1.3))
    safe_add_image(slide, "images/jacket-grey.jpg", Inches(4.3), Inches(3.4), Inches(1.3), Inches(1.3))

    tb_cap = slide.shapes.add_textbox(Inches(1.0), Inches(4.8), Inches(5.2), Inches(0.3))
    p_cap = tb_cap.text_frame.paragraphs[0]
    p_cap.text = "Generic catalog bestsellers: Running sneakers, Graphic tee, Denim jacket"
    p_cap.font.size = Pt(8.5)
    p_cap.font.color.rgb = COLOR_MUTED
    p_cap.alignment = PP_ALIGN.CENTER

    flaws_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.15), Inches(5.2), Inches(1.5))
    tf_fl = flaws_box.text_frame
    flaws = [
        ("❌ No size or stock check", "Fails to confirm if abandoned size is available."),
        ("❌ Completely misses intent", "Zero awareness that the shopper is shopping for a summer wedding."),
        ("❌ Ignores search history", "Recommends athletic sneakers instead of coordinating linen trousers.")
    ]
    for i, (f_title, f_desc) in enumerate(flaws):
        p_fl = tf_fl.paragraphs[0] if i == 0 else tf_fl.add_paragraph()
        if i > 0: p_fl.space_before = Pt(4)
        r_f1 = p_fl.add_run()
        r_f1.text = f"{f_title}: "
        r_f1.font.bold = True
        r_f1.font.size = Pt(9.5)
        r_f1.font.color.rgb = COLOR_RED
        r_f2 = p_fl.add_run()
        r_f2.text = f_desc
        r_f2.font.size = Pt(9.5)
        r_f2.font.color.rgb = COLOR_MUTED

    right_cards = [
        (
            "1. On-Site Search Knows. Email Doesn't.",
            "On-site search and product discovery sit in one tool; marketing campaigns sit in another. Each vendor sees only 50% of the customer's actual behavior.",
            "🧩 FRAGMENTED SILOS"
        ),
        (
            "2. Connecting Them Is a Project, Not a Setting.",
            "Bolt-on connectors pass a few delayed events across. Everything else the shopper searched, filtered, and compared stays trapped behind engineering backlogs.",
            "🔌 INTEGRATION FRICTION"
        ),
        (
            "3. AI Agents Only Act on What They Can See.",
            "Split the data and every AI agent works from half the story. As you add more specialized agents, the context gap grows exponentially.",
            "🤖 BLIND AI AGENTS"
        )
    ]

    card_y = 2.0
    for title, desc, tag in right_cards:
        rc = add_card(slide, Inches(6.8), Inches(card_y), Inches(5.7), Inches(1.5), COLOR_WHITE, COLOR_BORDER_LIGHT)
        tb_rc = slide.shapes.add_textbox(Inches(7.0), Inches(card_y + 0.1), Inches(5.3), Inches(1.3))
        tf_rc = tb_rc.text_frame
        tf_rc.word_wrap = True

        p_t = tf_rc.paragraphs[0]
        p_t.text = tag
        p_t.font.size = Pt(8.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_BLAZE

        p_h = tf_rc.add_paragraph()
        p_h.text = title
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_INK
        p_h.space_before = Pt(2)

        p_d = tf_rc.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_MUTED
        p_d.space_before = Pt(3)

        card_y += 1.7

def build_slide_3_journey(prs):
    """Slide 3: The Shopper Journey - Intent to Outbound Action (Dark Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_BASTILLE)

    add_header(
        slide,
        tag_text="Real-Time Intent Loop",
        title_text="What Shoppers Do on Your Site Writes Your Next Email",
        subtitle_text="Follow every shopper from search bar to inbox. Automatically capture intent in real time to build the next message.",
        dark_mode=True
    )

    steps = [
        ("2:02 PM", "Searches", "“linen shirt for a summer wedding”", "High-intent category & occasion captured"),
        ("2:05 PM", "Filters", "Sky blue, Size M, $60 to $90", "Specific style, sizing & budget locked"),
        ("2:09 PM", "Compares", "Opens Sky Blue Linen twice", "Deep product affinity identified"),
        ("2:14 PM", "Cart Exit", "Sky Blue Linen ($79) left in cart", "Immediate high-intent drop-off logged"),
        ("5:14 PM", "Triggered", "Automated email + WhatsApp sent", "Contextual recovery + matching chinos")
    ]

    step_w = 2.15
    start_x = 0.8
    for i, (time_str, action, detail, impact) in enumerate(steps):
        is_final = (i == 4)
        c_bg = COLOR_NIGHT if is_final else COLOR_CARD_DARK
        c_border = COLOR_BLAZE if is_final else COLOR_BORDER_DARK
        
        card = add_card(slide, Inches(start_x + i * 2.38), Inches(2.0), Inches(step_w), Inches(2.3), c_bg, c_border, 1.5 if is_final else 1)
        tb_s = slide.shapes.add_textbox(Inches(start_x + i * 2.38 + 0.12), Inches(2.1), Inches(step_w - 0.24), Inches(2.1))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True

        p_num = tf_s.paragraphs[0]
        p_num.text = f"STEP 0{i+1} • {time_str}"
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_BLAZE if is_final else COLOR_FAINT

        p_act = tf_s.add_paragraph()
        p_act.text = action
        p_act.font.size = Pt(12)
        p_act.font.bold = True
        p_act.font.color.rgb = COLOR_WHITE
        p_act.space_before = Pt(3)

        p_det = tf_s.add_paragraph()
        p_det.text = detail
        p_det.font.size = Pt(9.5)
        p_det.font.color.rgb = COLOR_BLAZE_LIGHT if is_final else COLOR_PEACH
        p_det.space_before = Pt(3)

        p_imp = tf_s.add_paragraph()
        p_imp.text = impact
        p_imp.font.size = Pt(8.5)
        p_imp.font.color.rgb = COLOR_FAINT
        p_imp.space_before = Pt(6)

    lp = add_card(slide, Inches(0.8), Inches(4.55), Inches(5.6), Inches(2.4), COLOR_CARD_DARK, COLOR_BORDER_DARK)
    tb_lp = slide.shapes.add_textbox(Inches(1.0), Inches(4.65), Inches(5.2), Inches(2.2))
    tf_lp = tb_lp.text_frame
    tf_lp.word_wrap = True

    p_lph = tf_lp.paragraphs[0]
    r_lph = p_lph.add_run()
    r_lph.text = "NETCORE REAL-TIME SHOPPER PROFILE "
    r_lph.font.size = Pt(9.5)
    r_lph.font.bold = True
    r_lph.font.color.rgb = COLOR_BLAZE
    r_lpb = p_lph.add_run()
    r_lpb.text = "[AUTOMATICALLY ENRICHED]"
    r_lpb.font.size = Pt(8.5)
    r_lpb.font.bold = True
    r_lpb.font.color.rgb = COLOR_GREEN

    log_rows = [
        ("2:02 PM", "Intent: summer wedding | Fabric: linen | Category: men's shirts"),
        ("2:05 PM", "Color: sky blue | Size: M | Budget ceiling: $90"),
        ("2:09 PM", "Top product: Sky blue linen shirt (viewed twice)"),
        ("2:14 PM", "Cart: $79 | Left without checkout"),
        ("5:14 PM", "Action: Send cart recovery email + matching chinos recommendation")
    ]
    for tm, desc in log_rows:
        p_r = tf_lp.add_paragraph()
        p_r.space_before = Pt(3)
        r_tm = p_r.add_run()
        r_tm.text = f"{tm} — "
        r_tm.font.bold = True
        r_tm.font.size = Pt(9)
        r_tm.font.color.rgb = COLOR_FAINT
        r_ds = p_r.add_run()
        r_ds.text = desc
        r_ds.font.size = Pt(9)
        r_ds.font.color.rgb = COLOR_WHITE

    rp = add_card(slide, Inches(6.8), Inches(4.55), Inches(5.7), Inches(2.4), COLOR_NIGHT, COLOR_BLAZE, 1.5)
    tb_rp = slide.shapes.add_textbox(Inches(7.0), Inches(4.65), Inches(5.3), Inches(2.2))
    tf_rp = tb_rp.text_frame
    tf_rp.word_wrap = True

    p_rph = tf_rp.paragraphs[0]
    p_rph.text = "THE NETCORE CONTEXT-AWARE MESSAGE (INBOX VIEW)"
    p_rph.font.size = Pt(9.5)
    p_rph.font.bold = True
    p_rph.font.color.rgb = COLOR_GREEN

    p_rps = tf_rp.add_paragraph()
    p_rps.text = "“Your sky blue linen shirt is still in your cart”"
    p_rps.font.size = Pt(13)
    p_rps.font.bold = True
    p_rps.font.color.rgb = COLOR_WHITE
    p_rps.space_before = Pt(3)

    p_rpb = tf_rp.add_paragraph()
    p_rpb.text = "“Size M is still in stock. Planning for the wedding? These pair well with it.”"
    p_rpb.font.size = Pt(9.5)
    p_rpb.font.color.rgb = COLOR_PEACH
    p_rpb.space_before = Pt(2)

    safe_add_image(slide, "images/shirt-sky.jpg", Inches(7.0), Inches(5.65), Inches(1.1), Inches(1.1))
    safe_add_image(slide, "images/pants-stone.jpg", Inches(9.4), Inches(5.65), Inches(1.1), Inches(1.1))

    tb_p1 = slide.shapes.add_textbox(Inches(8.15), Inches(5.75), Inches(1.2), Inches(0.9))
    p_p1 = tb_p1.text_frame.paragraphs[0]
    p_p1.text = "Sky Blue Linen\n$79 (In Cart)\nSize M"
    p_p1.font.size = Pt(8.5)
    p_p1.font.bold = True
    p_p1.font.color.rgb = COLOR_WHITE

    tb_p2 = slide.shapes.add_textbox(Inches(10.55), Inches(5.75), Inches(1.8), Inches(0.9))
    p_p2 = tb_p2.text_frame.paragraphs[0]
    p_p2.text = "Stone Linen Chinos\n$68 (Recommended)\nSize M"
    p_p2.font.size = Pt(8.5)
    p_p2.font.bold = True
    p_p2.font.color.rgb = COLOR_BLAZE_LIGHT

def build_slide_4_campaigns(prs):
    """Slide 4: Four High-Performance Behavioral Campaigns (Light Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_MIST)

    add_header(
        slide,
        tag_text="Omnichannel Campaign Playbook",
        title_text="Four Campaigns Built From What Shoppers Actually Do",
        subtitle_text="Combining on-site search intelligence with customer lifecycle marketing enables campaigns point solutions can never replicate.",
        dark_mode=False
    )

    # 4 Bento Cards Grid (2x2)
    # Top-Left: Search Segments
    c1 = add_card(slide, Inches(0.8), Inches(1.95), Inches(5.7), Inches(2.4), COLOR_WHITE, COLOR_BORDER_LIGHT)
    tb_c1 = slide.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(5.3), Inches(2.2))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True

    p1_h = tf_c1.paragraphs[0]
    p1_h.text = "CAMPAIGN 01"
    p1_h.font.size = Pt(9)
    p1_h.font.bold = True
    p1_h.font.color.rgb = COLOR_BLAZE

    p1_t = tf_c1.add_paragraph()
    p1_t.text = "Segments Built From Search Queries"
    p1_t.font.size = Pt(13)
    p1_t.font.bold = True
    p1_t.font.color.rgb = COLOR_INK
    p1_t.space_before = Pt(2)

    p1_d = tf_c1.add_paragraph()
    p1_d.text = "Reach shoppers by what they are looking for right now, not just what they bought last year."
    p1_d.font.size = Pt(9.5)
    p1_d.font.color.rgb = COLOR_MUTED
    p1_d.space_before = Pt(3)

    p1_ex = tf_c1.add_paragraph()
    p1_ex.text = "• Rule: Searched “linen” 2+ times  AND  Last 7 days, no purchase  AND  Size M\n• Dynamic Audience: 1,240 active shoppers ready for immediate conversion"
    p1_ex.font.size = Pt(9)
    p1_ex.font.bold = True
    p1_ex.font.color.rgb = COLOR_BLAZE
    p1_ex.space_before = Pt(6)

    # Top-Right: Matching Recommendations
    c2 = add_card(slide, Inches(6.8), Inches(1.95), Inches(5.7), Inches(2.4), COLOR_WHITE, COLOR_BORDER_LIGHT)
    tb_c2 = slide.shapes.add_textbox(Inches(7.0), Inches(2.05), Inches(5.3), Inches(2.2))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True

    p2_h = tf_c2.paragraphs[0]
    p2_h.text = "CAMPAIGN 02"
    p2_h.font.size = Pt(9)
    p2_h.font.bold = True
    p2_h.font.color.rgb = COLOR_BLAZE

    p2_t = tf_c2.add_paragraph()
    p2_t.text = "Email Recommendations That Match Your Site"
    p2_t.font.size = Pt(13)
    p2_t.font.bold = True
    p2_t.font.color.rgb = COLOR_INK
    p2_t.space_before = Pt(2)

    p2_d = tf_c2.add_paragraph()
    p2_d.text = "The exact same AI that ranks on-site search chooses products in every email, automatically respecting size and price."
    p2_d.font.size = Pt(9.5)
    p2_d.font.color.rgb = COLOR_MUTED
    p2_d.space_before = Pt(3)

    safe_add_image(slide, "images/shirt-sky.jpg", Inches(7.0), Inches(3.2), Inches(0.95), Inches(0.95))
    safe_add_image(slide, "images/pants-stone.jpg", Inches(8.3), Inches(3.2), Inches(0.95), Inches(0.95))
    safe_add_image(slide, "images/shoe-sand.jpg", Inches(9.6), Inches(3.2), Inches(0.95), Inches(0.95))
    safe_add_image(slide, "images/jacket-navy.jpg", Inches(10.9), Inches(3.2), Inches(0.95), Inches(0.95))

    # Bottom-Left: Merchandising Triggers
    c3 = add_card(slide, Inches(0.8), Inches(4.55), Inches(5.7), Inches(2.4), COLOR_WHITE, COLOR_BORDER_LIGHT)
    tb_c3 = slide.shapes.add_textbox(Inches(1.0), Inches(4.65), Inches(5.3), Inches(2.2))
    tf_c3 = tb_c3.text_frame
    tf_c3.word_wrap = True

    p3_h = tf_c3.paragraphs[0]
    p3_h.text = "CAMPAIGN 03"
    p3_h.font.size = Pt(9)
    p3_h.font.bold = True
    p3_h.font.color.rgb = COLOR_BLAZE

    p3_t = tf_c3.add_paragraph()
    p3_t.text = "Merchandising Triggers for the Right Shoppers"
    p3_t.font.size = Pt(13)
    p3_t.font.bold = True
    p3_t.font.color.rgb = COLOR_INK
    p3_t.space_before = Pt(2)

    p3_d = tf_c3.add_paragraph()
    p3_d.text = "Price drops, back-in-stock, and new arrivals message ONLY the shoppers who previously searched or viewed those products."
    p3_d.font.size = Pt(9.5)
    p3_d.font.color.rgb = COLOR_MUTED
    p3_d.space_before = Pt(3)

    p3_flow = tf_c3.add_paragraph()
    p3_flow.text = "⚡ WHEN: Linen shirt price drops to $69\n👥 WHO: Shoppers who searched or viewed it in last 14 days\n📲 SEND: Multi-channel blast via WhatsApp, SMS, and Email"
    p3_flow.font.size = Pt(9)
    p3_flow.font.bold = True
    p3_flow.font.color.rgb = COLOR_GREEN
    p3_flow.space_before = Pt(6)

    # Bottom-Right: Omnichannel Orbit (Adjusted width so text does not collide with image)
    c4 = add_card(slide, Inches(6.8), Inches(4.55), Inches(5.7), Inches(2.4), COLOR_WHITE, COLOR_BORDER_LIGHT)
    tb_c4 = slide.shapes.add_textbox(Inches(7.0), Inches(4.65), Inches(4.2), Inches(2.2))
    tf_c4 = tb_c4.text_frame
    tf_c4.word_wrap = True

    p4_h = tf_c4.paragraphs[0]
    p4_h.text = "CAMPAIGN 04"
    p4_h.font.size = Pt(9)
    p4_h.font.bold = True
    p4_h.font.color.rgb = COLOR_BLAZE

    p4_t = tf_c4.add_paragraph()
    p4_t.text = "Recommendations Across Every Channel"
    p4_t.font.size = Pt(13)
    p4_t.font.bold = True
    p4_t.font.color.rgb = COLOR_INK
    p4_t.space_before = Pt(2)

    p4_d = tf_c4.add_paragraph()
    p4_d.text = "Click a product in an email and the site leads with it. Browse on site, and the next push notification continues right where they left off."
    p4_d.font.size = Pt(9.5)
    p4_d.font.color.rgb = COLOR_MUTED
    p4_d.space_before = Pt(3)

    p4_chans = tf_c4.add_paragraph()
    p4_chans.text = "• Channels: Web • Email • SMS • Push • WhatsApp"
    p4_chans.font.size = Pt(9.5)
    p4_chans.font.bold = True
    p4_chans.font.color.rgb = COLOR_INK
    p4_chans.space_before = Pt(8)

    safe_add_image(slide, "images/shirt-sky.jpg", Inches(11.35), Inches(4.75), Inches(0.95), Inches(0.95))

def build_slide_5_memory_agents(prs):
    """Slide 5: Unified Shopper Memory & AI Agent Ecosystem (Dark Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_BASTILLE)

    add_header(
        slide,
        tag_text="Enterprise Architecture",
        title_text="One Shopper Memory Behind Every Experience",
        subtitle_text="Product discovery and customer engagement read from and write to the same profile. Nothing to sync, nothing lost.",
        dark_mode=True
    )

    col_w = 3.65
    gap = 0.38
    start_x = 0.8

    # Col 1: Inputs
    c_in = add_card(slide, Inches(start_x), Inches(1.95), Inches(col_w), Inches(2.75), COLOR_CARD_DARK, COLOR_BORDER_DARK)
    tb_in = slide.shapes.add_textbox(Inches(start_x + 0.15), Inches(2.05), Inches(col_w - 0.3), Inches(2.55))
    tf_in = tb_in.text_frame
    tf_in.word_wrap = True

    p_in_h = tf_in.paragraphs[0]
    p_in_h.text = "WHAT SHOPPERS DO (SIGNALS)"
    p_in_h.font.size = Pt(9.5)
    p_in_h.font.bold = True
    p_in_h.font.color.rgb = COLOR_FAINT

    signals = [
        "🔍 Searches & faceted filter refinements",
        "🖱️ Products viewed, compared & clicked",
        "🛒 Cart additions, checkouts & returns",
        "💬 Email, SMS & WhatsApp responses",
        "🏷️ Catalog sync (price drops, stock levels)",
        "📍 Context (device, geolocation, season)"
    ]
    for s in signals:
        p_s = tf_in.add_paragraph()
        p_s.text = s
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = COLOR_WHITE
        p_s.space_before = Pt(4)

    # Col 2: Central Memory Core
    c_core = add_card(slide, Inches(start_x + col_w + gap), Inches(1.95), Inches(col_w), Inches(2.75), COLOR_NIGHT, COLOR_BLAZE, 1.5)
    tb_core = slide.shapes.add_textbox(Inches(start_x + col_w + gap + 0.15), Inches(2.05), Inches(col_w - 0.3), Inches(2.55))
    tf_core = tb_core.text_frame
    tf_core.word_wrap = True

    p_ch = tf_core.paragraphs[0]
    p_ch.text = "NETCORE SHOPPER MEMORY"
    p_ch.font.size = Pt(9.5)
    p_ch.font.bold = True
    p_ch.font.color.rgb = COLOR_BLAZE

    p_ct = tf_core.add_paragraph()
    p_ct.text = "Real-Time Unified State Machine"
    p_ct.font.size = Pt(13)
    p_ct.font.bold = True
    p_ct.font.color.rgb = COLOR_WHITE
    p_ct.space_before = Pt(2)

    modes = [
        ("Returning Shoppers", "Builds upon historic LTV, past orders & cross-channel brand affinities."),
        ("New Shoppers", "Seeds intent from location, referring traffic, and catalog trending velocity."),
        ("In-Session Adaptation", "Dynamically updates recommendations and cart triggers with every click.")
    ]
    for m_t, m_d in modes:
        p_m = tf_core.add_paragraph()
        p_m.space_before = Pt(4)
        rm_t = p_m.add_run()
        rm_t.text = f"• {m_t}: "
        rm_t.font.bold = True
        rm_t.font.size = Pt(9)
        rm_t.font.color.rgb = COLOR_BLAZE_LIGHT
        rm_d = p_m.add_run()
        rm_d.text = m_d
        rm_d.font.size = Pt(9)
        rm_d.font.color.rgb = COLOR_FAINT

    # Col 3: Outbound & On-Site Surfaces
    c_out = add_card(slide, Inches(start_x + 2 * (col_w + gap)), Inches(1.95), Inches(col_w), Inches(2.75), COLOR_CARD_DARK, COLOR_BORDER_DARK)
    tb_out = slide.shapes.add_textbox(Inches(start_x + 2 * (col_w + gap) + 0.15), Inches(2.05), Inches(col_w - 0.3), Inches(2.55))
    tf_out = tb_out.text_frame
    tf_out.word_wrap = True

    p_out_h = tf_out.paragraphs[0]
    p_out_h.text = "WHERE IT SHOWS UP (SURFACES)"
    p_out_h.font.size = Pt(9.5)
    p_out_h.font.bold = True
    p_out_h.font.color.rgb = COLOR_FAINT

    p_onsite = tf_out.add_paragraph()
    p_onsite.text = "On Your Store:"
    p_onsite.font.bold = True
    p_onsite.font.size = Pt(10)
    p_onsite.font.color.rgb = COLOR_BLAZE_LIGHT
    p_onsite.space_before = Pt(4)

    p_onsite_items = tf_out.add_paragraph()
    p_onsite_items.text = "• On-site search ranking & visual auto-suggest\n• 1:1 Personalized product recommendations\n• Real-time visual merchandising badges"
    p_onsite_items.font.size = Pt(9)
    p_onsite_items.font.color.rgb = COLOR_WHITE
    p_onsite_items.space_before = Pt(2)

    p_camp = tf_out.add_paragraph()
    p_camp.text = "In Your Campaigns:"
    p_camp.font.bold = True
    p_camp.font.size = Pt(10)
    p_camp.font.color.rgb = COLOR_GREEN
    p_camp.space_before = Pt(6)

    p_camp_items = tf_out.add_paragraph()
    p_camp_items.text = "• Dynamic triggered & lifecycle emails\n• SMS and RCS conversational messaging\n• Mobile app push notifications & WhatsApp commerce"
    p_camp_items.font.size = Pt(9)
    p_camp_items.font.color.rgb = COLOR_WHITE
    p_camp_items.space_before = Pt(2)

    # Bottom Full Card: 9 Autonomous AI Agents
    c_agents = add_card(slide, Inches(0.8), Inches(4.9), Inches(11.7), Inches(2.1), COLOR_CARD_DARK, COLOR_BORDER_DARK)
    tb_ag = slide.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.9))
    tf_ag = tb_ag.text_frame
    tf_ag.word_wrap = True

    p_agh = tf_ag.paragraphs[0]
    p_agh.text = "9 SPECIALIZED AUTONOMOUS AI AGENTS WORKING FROM ONE MEMORY"
    p_agh.font.size = Pt(9.5)
    p_agh.font.bold = True
    p_agh.font.color.rgb = COLOR_BLAZE

    p_on_ag = tf_ag.add_paragraph()
    p_on_ag.space_before = Pt(4)
    r_on_t = p_on_ag.add_run()
    r_on_t.text = "On Your Site:   "
    r_on_t.font.bold = True
    r_on_t.font.size = Pt(10)
    r_on_t.font.color.rgb = COLOR_WHITE
    r_on_b = p_on_ag.add_run()
    r_on_b.text = "🛍️ Shopping Agent   •   🗂️ Merchandising Agent   •   🛠️ Debugger Agent"
    r_on_b.font.size = Pt(10)
    r_on_b.font.color.rgb = COLOR_BLAZE_LIGHT

    p_cmp_ag = tf_ag.add_paragraph()
    p_cmp_ag.space_before = Pt(4)
    r_cmp_t = p_cmp_ag.add_run()
    r_cmp_t.text = "In Campaigns:   "
    r_cmp_t.font.bold = True
    r_cmp_t.font.size = Pt(10)
    r_cmp_t.font.color.rgb = COLOR_WHITE
    r_cmp_b = p_cmp_ag.add_run()
    r_cmp_b.text = "👥 Audience Agent  •  ✍️ Content Agent  •  📅 Scheduler Agent  •  🔀 Decisioning Agent  •  📊 Insights Agent  •  🚀 Co-Marketer"
    r_cmp_b.font.size = Pt(10)
    r_cmp_b.font.color.rgb = COLOR_GREEN

    p_note = tf_ag.add_paragraph()
    p_note.text = "Every agent operates on 100% complete customer context — no data blindness, no cross-tool latency."
    p_note.font.size = Pt(8.5)
    p_note.font.color.rgb = COLOR_FAINT
    p_note.space_before = Pt(6)

def build_slide_6_transformation(prs):
    """Slide 6: Operational Transformation for Marketing Teams (Light Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_MIST)

    add_header(
        slide,
        tag_text="Business Impact & Consolidation",
        title_text="What Changes for Your Marketing Team",
        subtitle_text="Replace four fragmented point solutions with one unified platform. Less time stitching tools together, more campaigns that sell.",
        dark_mode=False
    )

    # Left Box: The Fragmented 4-Vendor Reality
    c_left = add_card(slide, Inches(0.8), Inches(1.95), Inches(5.6), Inches(4.9), COLOR_WHITE, COLOR_RED, border_width=1.5)
    tb_tl = slide.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.2), Inches(4.6))
    tf_tl = tb_tl.text_frame
    tf_tl.word_wrap = True

    p_lh = tf_tl.paragraphs[0]
    p_lh.text = "THE LEGACY FRAGMENTED STACK"
    p_lh.font.size = Pt(9.5)
    p_lh.font.bold = True
    p_lh.font.color.rgb = COLOR_RED

    p_sub_l = tf_tl.add_paragraph()
    p_sub_l.text = "Four Disjointed Tools, Four Invoices"
    p_sub_l.font.size = Pt(15)
    p_sub_l.font.bold = True
    p_sub_l.font.color.rgb = COLOR_INK
    p_sub_l.space_before = Pt(3)

    invoices = [
        ("1. On-Site Search & Discovery Vendor", "Separate contract, custom event tracking, catalog sync latency."),
        ("2. Standalone ESP / Cross-Channel Vendor", "Disconnected profile data, limited browsing events, costly volume tiers."),
        ("3. Third-Party Personalization Engine", "Custom tag management, fragile script execution, inconsistent rules."),
        ("4. Dedicated Deliverability Tool", "Separate monitoring dashboard, vendor finger-pointing when issues arise.")
    ]
    for inv_title, inv_sub in invoices:
        p_inv = tf_tl.add_paragraph()
        p_inv.space_before = Pt(8)
        r_it = p_inv.add_run()
        r_it.text = f"📄 {inv_title}\n"
        r_it.font.bold = True
        r_it.font.size = Pt(10.5)
        r_it.font.color.rgb = COLOR_INK
        r_is = p_inv.add_run()
        r_is.text = f"    {inv_sub}"
        r_is.font.size = Pt(9)
        r_is.font.color.rgb = COLOR_MUTED

    p_cost = tf_tl.add_paragraph()
    p_cost.text = "Result: High engineering overhead, fragmented data, delayed campaigns."
    p_cost.font.size = Pt(9.5)
    p_cost.font.bold = True
    p_cost.font.color.rgb = COLOR_RED
    p_cost.space_before = Pt(14)

    # Right Column: The 4 Strategic Pillars of Netcore
    c_right = add_card(slide, Inches(6.8), Inches(1.95), Inches(5.7), Inches(4.9), COLOR_WHITE, COLOR_BLAZE, border_width=1.5)
    tb_tr = slide.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.6))
    tf_tr = tb_tr.text_frame
    tf_tr.word_wrap = True

    p_rh = tf_tr.paragraphs[0]
    p_rh.text = "THE NETCORE ADVANTAGE"
    p_rh.font.size = Pt(9.5)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_BLAZE

    p_sub_r = tf_tr.add_paragraph()
    p_sub_r.text = "One Platform • One Invoice • One Partner"
    p_sub_r.font.size = Pt(15)
    p_sub_r.font.bold = True
    p_sub_r.font.color.rgb = COLOR_INK
    p_sub_r.space_before = Pt(3)

    pillars = [
        ("🚀 More Revenue From Every Send", "Messages built on real-time shopping intent get opened, clicked, and converted. New York & Company saw 70% higher revenue per email."),
        ("⚡ Zero-Integration Campaigns", "Search, browse, cart, and campaign data live together natively. No custom syncs, no event mapping, no developer backlogs."),
        ("📋 Unified Procurement & Support", "Consolidate 4 vendor contracts, renewals, and support tickets into a single enterprise SLA and trusted account team."),
        ("🧠 Smarter, Context-Rich AI", "Every AI agent operates on the complete shopper memory, getting sharper with every search query and catalog update.")
    ]
    for pil_t, pil_d in pillars:
        p_pil = tf_tr.add_paragraph()
        p_pil.space_before = Pt(8)
        rp_t = p_pil.add_run()
        rp_t.text = f"{pil_t}\n"
        rp_t.font.bold = True
        rp_t.font.size = Pt(10.5)
        rp_t.font.color.rgb = COLOR_INK
        rp_d = p_pil.add_run()
        rp_d.text = pil_d
        rp_d.font.size = Pt(9)
        rp_d.font.color.rgb = COLOR_MUTED

def build_slide_7_results(prs):
    """Slide 7: Enterprise Case Study & Industry Analyst Recognition (Dark Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_BASTILLE)

    add_header(
        slide,
        tag_text="Proven Enterprise Scale",
        title_text="Real Results From Brands Running Both Together",
        subtitle_text="Validated by global commerce leaders and top-tier independent industry analyst firms.",
        dark_mode=True
    )

    # Left Half: Featured Case Study (New York & Company)
    cs_card = add_card(slide, Inches(0.8), Inches(1.95), Inches(5.6), Inches(4.9), COLOR_CARD_DARK, COLOR_BORDER_DARK)
    tb_cs = slide.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.2), Inches(4.5))
    tf_cs = tb_cs.text_frame
    tf_cs.word_wrap = True

    p_csh = tf_cs.paragraphs[0]
    p_csh.text = "FEATURED CASE STUDY: NEW YORK & COMPANY"
    p_csh.font.size = Pt(9.5)
    p_csh.font.bold = True
    p_csh.font.color.rgb = COLOR_BLAZE

    p_css = tf_cs.add_paragraph()
    p_css.text = "Unifying On-Site Search & Personalized Lifecycle"
    p_css.font.size = Pt(13)
    p_css.font.bold = True
    p_css.font.color.rgb = COLOR_WHITE
    p_css.space_before = Pt(2)

    p_csd = tf_cs.add_paragraph()
    p_csd.text = "On-site search learned from live shopper behavior, and those exact signals instantly triggered personalized emails and SMS."
    p_csd.font.size = Pt(9.5)
    p_csd.font.color.rgb = COLOR_FAINT
    p_csd.space_before = Pt(3)

    stat_data = [
        ("+600%", "Higher Email Click Rate"),
        ("+70%", "More Revenue Per Email"),
        ("+8%", "Higher Search Conversion"),
        ("1,228%", "Verified Return on Investment (ROI)")
    ]
    for i, (num, lbl) in enumerate(stat_data):
        rx = 1.0 + (i % 2) * 2.65
        ry = 3.35 + (i // 2) * 1.05
        sc = add_card(slide, Inches(rx), Inches(ry), Inches(2.45), Inches(0.95), COLOR_NIGHT, COLOR_BORDER_DARK)
        tb_sc = slide.shapes.add_textbox(Inches(rx + 0.1), Inches(ry + 0.08), Inches(2.25), Inches(0.8))
        tf_sc = tb_sc.text_frame
        p_n = tf_sc.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(18)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_BLAZE
        p_l = tf_sc.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(8.5)
        p_l.font.color.rgb = COLOR_WHITE
        p_l.space_before = Pt(1)

    q_card = add_card(slide, Inches(1.0), Inches(5.6), Inches(5.2), Inches(1.05), COLOR_BASTILLE, COLOR_BLAZE, 1)
    tb_q = slide.shapes.add_textbox(Inches(1.15), Inches(5.65), Inches(4.9), Inches(0.95))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True
    p_qt = tf_q.paragraphs[0]
    p_qt.text = "“With AI-powered search and personalized emails working together, we created a smoother shopping journey. Every click fueled smarter marketing.”"
    p_qt.font.size = Pt(8.5)
    p_qt.font.italic = True
    p_qt.font.color.rgb = COLOR_PEACH
    p_qc = tf_q.add_paragraph()
    p_qc.text = "— VP of Marketing, New York & Company"
    p_qc.font.size = Pt(8)
    p_qc.font.bold = True
    p_qc.font.color.rgb = COLOR_WHITE
    p_qc.space_before = Pt(2)

    # Right Half: Analyst Recognition
    an_card = add_card(slide, Inches(6.8), Inches(1.95), Inches(5.7), Inches(4.9), COLOR_CARD_DARK, COLOR_BORDER_DARK)
    tb_an = slide.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.5))
    tf_an = tb_an.text_frame
    tf_an.word_wrap = True

    p_anh = tf_an.paragraphs[0]
    p_anh.text = "INDUSTRY ANALYST VALIDATION"
    p_anh.font.size = Pt(9.5)
    p_anh.font.bold = True
    p_anh.font.color.rgb = COLOR_GREEN

    p_ant = tf_an.add_paragraph()
    p_ant.text = "Recognized on Both Sides of the Stack"
    p_ant.font.size = Pt(13)
    p_ant.font.bold = True
    p_ant.font.color.rgb = COLOR_WHITE
    p_ant.space_before = Pt(2)

    analysts = [
        (
            "GARTNER MAGIC QUADRANT — LEADER",
            "Named a Leader in the 2026 Gartner® Magic Quadrant™ for Search and Product Discovery for the 3rd consecutive year.",
            COLOR_BLAZE
        ),
        (
            "FORRESTER WAVE — STRONG PERFORMER",
            "Recognized as a Strong Performer in The Forrester Wave™: Email Marketing Service Providers (Q1 2026).",
            COLOR_GREEN
        ),
        (
            "FORRESTER WAVE — COMMERCE SEARCH",
            "Achieved the Highest Score in Current Offering in The Forrester Wave™: Commerce Search & Product Discovery (Q3 2025).",
            COLOR_BLAZE_LIGHT
        )
    ]

    card_ay = 3.05
    for tag_txt, body_txt, accent_c in analysts:
        ac = add_card(slide, Inches(7.0), Inches(card_ay), Inches(5.3), Inches(1.05), COLOR_NIGHT, accent_c, 1)
        tb_a = slide.shapes.add_textbox(Inches(7.15), Inches(card_ay + 0.08), Inches(5.0), Inches(0.9))
        tf_a = tb_a.text_frame
        tf_a.word_wrap = True

        pa_t = tf_a.paragraphs[0]
        pa_t.text = tag_txt
        pa_t.font.size = Pt(9)
        pa_t.font.bold = True
        pa_t.font.color.rgb = accent_c

        pa_b = tf_a.add_paragraph()
        pa_b.text = body_txt
        pa_b.font.size = Pt(8.5)
        pa_b.font.color.rgb = COLOR_WHITE
        pa_b.space_before = Pt(2)

        card_ay += 1.2

def build_slide_8_conclusion(prs):
    """Slide 8: Executive Summary & Call to Action (Dark Bastille Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, prs, COLOR_BASTILLE)

    add_header(
        slide,
        tag_text="Summary & Next Steps",
        title_text="Your Shoppers Already Told You What They Want",
        subtitle_text="Put every search, click, and cart to work across the next email, SMS, push, and WhatsApp message you send.",
        dark_mode=True
    )

    pillars = [
        (
            "1. Unify Search & Lifecycle",
            "Eliminate data lag and architectural silos between what shoppers do on site and what your campaigns deliver.",
            "🔄 SEAMLESS INTEGRATION",
            COLOR_BLAZE
        ),
        (
            "2. Trigger Real-Time Intent",
            "Transition from generic batch-and-blast emails to automated 1:1 triggers within minutes of customer activity.",
            "⚡ IMMEDIATE ACTION",
            COLOR_GREEN
        ),
        (
            "3. Maximize Commerce GMV",
            "Drive double-digit revenue growth with search-driven recommendations and synchronized omnichannel touchpoints.",
            "📈 PROVEN ROI",
            COLOR_BLAZE_LIGHT
        )
    ]

    col_w = 3.65
    gap = 0.38
    start_x = 0.8
    for i, (title, desc, badge, col) in enumerate(pillars):
        c = add_card(slide, Inches(start_x + i * (col_w + gap)), Inches(2.1), Inches(col_w), Inches(2.3), COLOR_CARD_DARK, col, 1.5)
        tb = slide.shapes.add_textbox(Inches(start_x + i * (col_w + gap) + 0.15), Inches(2.25), Inches(col_w - 0.3), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p_b = tf.paragraphs[0]
        p_b.text = badge
        p_b.font.size = Pt(9)
        p_b.font.bold = True
        p_b.font.color.rgb = col

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(3)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_FAINT
        p_d.space_before = Pt(4)

    cta_card = add_card(slide, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.2), COLOR_NIGHT, COLOR_BLAZE, 2.0)
    tb_cta = slide.shapes.add_textbox(Inches(1.2), Inches(4.85), Inches(11.0), Inches(1.9))
    tf_cta = tb_cta.text_frame
    tf_cta.word_wrap = True

    p_ch = tf_cta.paragraphs[0]
    p_ch.text = "READY TO TRANSFORM YOUR COMMERCE EXPERIENCE?"
    p_ch.font.size = Pt(10)
    p_ch.font.bold = True
    p_ch.font.color.rgb = COLOR_BLAZE

    p_ct = tf_cta.add_paragraph()
    p_ct.text = "Book a Live Architectural Walkthrough & Product Discovery Audit"
    p_ct.font.size = Pt(17)
    p_ct.font.bold = True
    p_ct.font.color.rgb = COLOR_WHITE
    p_ct.space_before = Pt(2)

    p_cd = tf_cta.add_paragraph()
    p_cd.text = (
        "See how Netcore connects on-site search, merchandising, and cross-channel lifecycle marketing "
        "to deliver immediate double-digit conversion lift for your brand."
    )
    p_cd.font.size = Pt(10.5)
    p_cd.font.color.rgb = COLOR_PEACH
    p_cd.space_before = Pt(3)

    btn_pill = add_card(slide, Inches(1.2), Inches(6.15), Inches(3.6), Inches(0.48), COLOR_BLAZE, COLOR_BLAZE_LIGHT)
    p_btn = btn_pill.text_frame.paragraphs[0]
    p_btn.text = "BOOK A DEMO  •  netcorecloud.com"
    p_btn.font.name = FONT_FAMILY
    p_btn.font.size = Pt(10)
    p_btn.font.bold = True
    p_btn.font.color.rgb = COLOR_WHITE
    p_btn.alignment = PP_ALIGN.CENTER

    tb_cnt = slide.shapes.add_textbox(Inches(7.2), Inches(6.15), Inches(5.0), Inches(0.45))
    p_cnt = tb_cnt.text_frame.paragraphs[0]
    p_cnt.text = "Trusted by 6,500+ global brands  |  Gartner Magic Quadrant Leader"
    p_cnt.font.size = Pt(9.5)
    p_cnt.font.color.rgb = COLOR_FAINT
    p_cnt.alignment = PP_ALIGN.RIGHT

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    print("Building Slide 1: Hero & Connected Overview...")
    build_slide_1_hero(prs)

    print("Building Slide 2: The Core Problem (Disconnected Stacks)...")
    build_slide_2_problem(prs)

    print("Building Slide 3: The Shopper Journey...")
    build_slide_3_journey(prs)

    print("Building Slide 4: Four High-Performance Campaigns...")
    build_slide_4_campaigns(prs)

    print("Building Slide 5: Single Shopper Memory & AI Agents...")
    build_slide_5_memory_agents(prs)

    print("Building Slide 6: Operational Transformation for Marketing...")
    build_slide_6_transformation(prs)

    print("Building Slide 7: Enterprise Case Study & Recognition...")
    build_slide_7_results(prs)

    print("Building Slide 8: Executive Summary & Call to Action...")
    build_slide_8_conclusion(prs)

    output_path = "Netcore_Unified_Commerce_Deck.pptx"
    prs.save(output_path)
    print(f"\nPresentation successfully generated: {os.path.abspath(output_path)}")
    print(f"File size: {os.path.getsize(output_path):,} bytes")

if __name__ == "__main__":
    main()
